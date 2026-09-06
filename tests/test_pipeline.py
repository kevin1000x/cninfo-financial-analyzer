"""
Unit tests for pipeline output-quality safeguards.
"""

import json
import os
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.pipeline import (
    FinancialAnalysisPipeline,
    StreamingManifest,
    select_analysis_text,
)


class StubAnalyzer:
    """Capture analyze_text inputs while returning a minimal result payload."""

    def __init__(self):
        self.calls = []

    def analyze_text(self, text, file_path=None):
        self.calls.append((text, file_path))
        return {
            'tone_raw': 0.1,
            'fog_index': 3.2,
            'pos_word_count': 1,
            'neg_word_count': 0,
            'total_words': 10,
            'avg_sentence_length': 5.0,
            'complex_word_count': 2,
            'complex_word_pct': 20.0,
            'sentence_count': 2,
            'char_count': len(text),
            'file_size_bytes': 128,
            'success': True,
            'error': None,
        }


class StubStreamingDownloader:
    """Small downloader double that creates a local PDF placeholder."""

    def __init__(self, download_path):
        self.download_path = str(download_path)
        self.downloaded_paths = []

    def build_download_url(self, announcement):
        return 'https://example.com/report.pdf', announcement.get('filename', 'report.pdf')

    def query_announcements(self, stock_code, year, report_type):
        return [_streaming_announcement(len(self.downloaded_paths))]

    def download_one_sync(self, url, save_path):
        Path(save_path).parent.mkdir(parents=True, exist_ok=True)
        Path(save_path).write_bytes(b'%PDF-1.4\n% streaming test placeholder\n')
        self.downloaded_paths.append(save_path)
        return True


class StubStreamingParser:
    """Parser double that records save_output and returns generated parse results."""

    def __init__(self, result_factory):
        self.result_factory = result_factory
        self.calls = []

    def parse_pdf(self, pdf_path, save_output=True, extract_tables=None):
        self.calls.append((pdf_path, save_output, extract_tables))
        return self.result_factory(pdf_path, save_output)


@pytest.fixture
def pipeline_stub(tmp_path):
    """Create a lightweight pipeline instance without full initialization."""
    pipeline = FinancialAnalysisPipeline.__new__(FinancialAnalysisPipeline)
    pipeline.config = {
        'output': {
            'results_path': str(tmp_path),
            'format': 'excel',
            'summary_columns': [
                'stock_code',
                'company_name',
                'year',
                'report_type',
                'file_path',
                'text_path',
                'analysis_text_source',
                'download_date',
                'tone_raw',
                'fog_index',
            ],
        }
    }
    pipeline.analyzer = StubAnalyzer()
    pipeline.analysis_text_min_chars = 500
    pipeline.analysis_text_min_ratio = 0.02
    pipeline.save_intermediate = False
    pipeline.intermediate_output_path = str(tmp_path / 'intermediate')
    pipeline.max_saved_intermediates = 5
    pipeline.streaming_audit_output_path = str(tmp_path / 'streaming_audit')
    pipeline.streaming_max_audit_reports = 10
    pipeline.streaming_manifest_path = str(tmp_path / 'streaming_manifest.json')
    pipeline.company_map = {}
    return pipeline


@pytest.fixture
def manifest(tmp_path):
    """Fresh ledger for tests that drive _process_single_report directly."""
    return StreamingManifest(str(tmp_path / 'manifest.json'))


def _good_streaming_texts():
    mda_text = "第三章 管理层讨论与分析\n" + ("经营情况良好，营业收入持续增长。" * 100)
    full_text = mda_text + "\n" + ("这是完整报告正文。" * 600)
    return full_text, mda_text


def _streaming_announcement(idx=0):
    return {
        'announcementTitle': f'2023年年度报告_{idx}',
        'adjunctUrl': f'finalpage/2024-03-15/{idx}.PDF',
        'announcementTime': 1710432000000 + idx,
        'filename': f'report_{idx}.pdf',
    }


def test_analyze_phase_falls_back_to_full_text_for_bad_mda(pipeline_stub):
    """Short TOC-like MD&A text should not drive final metrics."""
    full_text = "第三章 管理层讨论与分析\n" + ("这是实际正文内容。" * 200)
    parse_results = pd.DataFrame([
        {
            'stock_code': '000001',
            'company_name': '平安银行',
            'year': 2023,
            'report_type': 'annual',
            'file_path': 'data/raw/report.pdf',
            'text_path': 'data/parsed/full_text.txt',
            'download_date': 1710432000000,
            'text': full_text,
            'mda_text': '管理层讨论与分析 ................................ 22\n3.1 总体经营情况 ........................... 22',
        }
    ])

    results = pipeline_stub.analyze_phase(parse_results)

    assert len(results) == 1
    assert pipeline_stub.analyzer.calls[0][0] == full_text
    assert results.iloc[0]['analysis_text_source'] == 'full_text'
    assert results.iloc[0]['download_date'] == '2024-03-15'


def test_analyze_phase_labels_parser_whole_document_fallback_as_full_text(
        pipeline_stub, tmp_path):
    """Cross-module guard: PDFParser.extract_mda_section returns the entire
    document when no candidate validates, and analyze_phase must not export
    that under the mda_text label."""
    from src.pdf_parser import PDFParser

    report_text = (
        "目 录\n"
        "第三章 管理层讨论与分析 ................................ 22\n"
        "第四章 公司治理 ......................................... 62\n"
    ) + ("本年度公司经营状况说明。" * 300)

    parser = PDFParser({'parser': {'output_path': str(tmp_path)}})
    mda_text = parser.extract_mda_section(report_text)
    assert mda_text == report_text, "fixture must exercise the parser fallback"

    results = pipeline_stub.analyze_phase(pd.DataFrame([
        {
            'stock_code': '000001',
            'company_name': '平安银行',
            'year': 2023,
            'report_type': 'annual',
            'file_path': 'data/raw/report.pdf',
            'text_path': 'data/parsed/full_text.txt',
            'download_date': 1710432000000,
            'text': report_text,
            'mda_text': mda_text,
        }
    ]))

    assert results.iloc[0]['analysis_text_source'] == 'full_text'
    assert pipeline_stub.analyzer.calls[0][0] == report_text


def test_select_analysis_text_uses_valid_mda_text():
    """Reusable text selector should prefer valid MD&A content."""
    mda_text = "第三章 管理层讨论与分析\n" + ("经营情况良好，收入持续增长。" * 80)
    full_text = "第一章 公司情况\n" + ("其他正文。" * 600) + "\n" + mda_text

    text, source = select_analysis_text({
        'text': full_text,
        'mda_text': mda_text,
    })

    assert text == mda_text
    assert source == 'mda_text'


def test_select_analysis_text_rejects_whole_document_fallback():
    """extract_mda_section returns the entire document when no candidate
    validates; that must never be scored or labelled as MD&A."""
    full_text = "第三章 管理层讨论与分析\n" + ("这是完整正文内容。" * 500)
    parse_result = {'text': full_text, 'mda_text': full_text}

    text, source = select_analysis_text(parse_result)

    assert source == 'full_text'
    assert text == full_text


def test_select_analysis_text_rejects_near_total_mda_slice():
    """A slice covering almost the whole report is the fallback, not a section."""
    body = "这是完整正文内容。" * 500
    mda_text = body[: int(len(body) * 0.9)]

    text, source = select_analysis_text({'text': body, 'mda_text': mda_text})

    assert source == 'full_text'
    assert text == body


def test_select_analysis_text_keeps_large_but_plausible_mda():
    """The upper guard must not reject a genuine MD&A from a short report."""
    mda_text = "第三章 管理层讨论与分析\n" + ("经营情况良好，收入持续增长。" * 100)
    full_text = mda_text + "\n" + ("其他正文。" * 600)

    text, source = select_analysis_text({'text': full_text, 'mda_text': mda_text})

    assert source == 'mda_text'
    assert text == mda_text


def test_select_analysis_text_falls_back_for_toc_mda():
    """TOC-like MD&A snippets should fall back to full text."""
    full_text = "第三章 管理层讨论与分析\n" + ("这是完整正文内容。" * 200)

    text, source = select_analysis_text({
        'text': full_text,
        'mda_text': '管理层讨论与分析 ................................ 22\n3.1 总体经营情况 ........................... 22',
    })

    assert text == full_text
    assert source == 'full_text'


def test_select_analysis_text_falls_back_for_short_mda():
    """Short non-TOC MD&A text should fall back when full text is available."""
    full_text = "第三章 管理层讨论与分析\n" + ("这是完整正文内容。" * 200)
    short_mda = "管理层讨论与分析。本年收入增长。"

    text, source = select_analysis_text({
        'text': full_text,
        'mda_text': short_mda,
    })

    assert text == full_text
    assert source == 'full_text'


def test_select_analysis_text_falls_back_for_low_mda_ratio():
    """Tiny MD&A slices relative to full text should fall back to full text."""
    full_text = "第三章 管理层讨论与分析\n" + ("这是完整正文内容。" * 5000)
    mda_text = "第三章 管理层讨论与分析\n" + ("经营情况良好。" * 80)

    text, source = select_analysis_text({
        'text': full_text,
        'mda_text': mda_text,
    })

    assert text == full_text
    assert source == 'full_text'


def test_pipeline_public_text_selector_uses_configured_thresholds(pipeline_stub):
    """External callers can reuse the pipeline-configured text selector."""
    pipeline_stub.analysis_text_min_chars = 20
    pipeline_stub.analysis_text_min_ratio = 0.5
    full_text = "第三章 管理层讨论与分析\n" + ("完整正文。" * 100)
    mda_text = "第三章 管理层讨论与分析\n" + ("经营良好。" * 10)

    text, source = pipeline_stub.select_analysis_text({
        'text': full_text,
        'mda_text': mda_text,
    })

    assert text == full_text
    assert source == 'full_text'


def test_save_results_normalizes_exported_types(pipeline_stub, monkeypatch):
    """Exported workbook should preserve padded stock codes and readable dates."""
    monkeypatch.setattr('src.pipeline.create_timestamp', lambda: '20240101_120000')
    results_df = pd.DataFrame([
        {
            'stock_code': 1,
            'company_name': '平安银行',
            'year': 2023,
            'report_type': 'annual',
            'file_path': 'data/raw/report.pdf',
            'text_path': 'data/parsed/full_text.txt',
            'analysis_text_source': 'full_text',
            'download_date': 1710432000000,
            'tone_raw': 0.1,
            'fog_index': 3.2,
        }
    ])

    output_path = pipeline_stub.save_results(results_df)

    assert os.path.exists(output_path)
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    types = [cell.data_type for cell in next(sheet.iter_rows(min_row=2, max_row=2))]

    stock_idx = headers.index('stock_code')
    date_idx = headers.index('download_date')
    assert values[stock_idx] == '000001'
    assert types[stock_idx] == 's'
    assert values[date_idx] == '2024-03-15'


def test_save_results_writes_schema_for_empty_dataframe(pipeline_stub, monkeypatch):
    """Even empty exports should keep a stable schema for downstream use."""
    monkeypatch.setattr('src.pipeline.create_timestamp', lambda: '20240101_120001')

    output_path = pipeline_stub.save_results(pd.DataFrame())

    exported = pd.read_excel(output_path)
    assert list(exported.columns) == pipeline_stub.config['output']['summary_columns']
    assert exported.empty


def test_streaming_audit_retention_keeps_latest_dirs(pipeline_stub, tmp_path):
    """The audit cap is an end-of-run step: relocating reports never prunes."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.streaming_max_audit_reports = 2

    results = []
    for idx in range(3):
        source_dir = tmp_path / f'tmp_parse_{idx}'
        source_dir.mkdir()
        (source_dir / 'full_text.txt').write_text(f'full-{idx}', encoding='utf-8')
        (source_dir / 'mda_text.txt').write_text(f'mda-{idx}', encoding='utf-8')

        parse_result = {
            'output_dir': str(source_dir),
            'text_path': str(source_dir / 'full_text.txt'),
            'mda_path': str(source_dir / 'mda_text.txt'),
        }

        relocated = pipeline_stub._relocate_streaming_audit_artifacts(
            parse_result=parse_result,
            stock_code=f'{idx + 1:06d}',
            year=2023,
            report_type='annual',
            announcement={
                'announcementTitle': f'2023年年度报告_{idx}',
                'adjunctUrl': f'finalpage/2024-03-15/{idx}.PDF',
                'announcementTime': 1710432000000 + idx,
            }
        )
        results.append(relocated)

        if idx == 0:
            oldest_dir = Path(relocated['output_dir'])
            os.utime(oldest_dir, (1, 1))
        elif idx == 1:
            middle_dir = Path(relocated['output_dir'])
            os.utime(middle_dir, (2, 2))

    assert all(Path(item['output_dir']).exists() for item in results), \
        "relocating a report must not prune earlier ones"

    pipeline_stub._enforce_streaming_audit_retention()

    audit_dirs = sorted(
        path.name for path in Path(pipeline_stub.streaming_audit_output_path).iterdir()
        if path.is_dir()
    )

    assert len(audit_dirs) == 2
    assert not Path(results[0]['output_dir']).exists()
    assert Path(results[1]['output_dir']).exists()
    assert Path(results[2]['output_dir']).exists()


def test_streaming_save_parsed_text_false_leaves_no_audit_dir(pipeline_stub, tmp_path, manifest):
    """Without parsed-text persistence, streaming should not leave audit artifacts."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    full_text, mda_text = _good_streaming_texts()

    pipeline_stub.parser = StubStreamingParser(
        lambda _pdf_path, _save_output: {
            'text': full_text,
            'mda_text': mda_text,
            'text_path': '',
            'mda_path': '',
            'output_dir': '',
        }
    )

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=False,
        manifest=manifest,
    )

    assert result is not None
    assert pipeline_stub.parser.calls[0][1] is False
    assert result['text_path'] == ''
    assert result['file_path'] == ''
    audit_root = Path(pipeline_stub.streaming_audit_output_path)
    assert not audit_root.exists() or not any(audit_root.iterdir())
    assert not Path(pipeline_stub.downloader.downloaded_paths[0]).exists()


def test_streaming_skips_table_extraction(pipeline_stub, tmp_path, manifest):
    """Streaming analysis never reads tables, so it must not pay to extract them.

    select_analysis_text only consumes text/mda_text; the table engine was
    most of the per-report parse time and its CSVs went unread."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    full_text, mda_text = _good_streaming_texts()

    pipeline_stub.parser = StubStreamingParser(
        lambda _pdf_path, _save_output: {
            'text': full_text,
            'mda_text': mda_text,
            'text_path': '',
            'mda_path': '',
            'output_dir': '',
        }
    )

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=False,
        manifest=manifest,
    )

    assert result is not None
    assert pipeline_stub.parser.calls[0][2] is False


def test_streaming_raw_text_disabled_keeps_audit_dir_without_txt(pipeline_stub, tmp_path, manifest):
    """If parser saves an audit directory without raw text, relocation should preserve that."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    full_text, mda_text = _good_streaming_texts()

    def parse_result(_pdf_path, _save_output):
        source_dir = tmp_path / 'parser_output_no_txt'
        source_dir.mkdir()
        (source_dir / 'table_1.csv').write_text('col\nvalue\n', encoding='utf-8')
        return {
            'text': full_text,
            'mda_text': mda_text,
            'text_path': '',
            'mda_path': '',
            'output_dir': str(source_dir),
        }

    pipeline_stub.parser = StubStreamingParser(parse_result)

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=True,
        manifest=manifest,
    )

    audit_dirs = [path for path in Path(pipeline_stub.streaming_audit_output_path).iterdir()]
    assert result is not None
    assert pipeline_stub.parser.calls[0][1] is True
    assert len(audit_dirs) == 1
    audit_dir = audit_dirs[0]
    assert audit_dir.is_dir()
    assert not (audit_dir / 'full_text.txt').exists()
    assert not (audit_dir / 'mda_text.txt').exists()
    assert (audit_dir / 'table_1.csv').exists()
    assert result['text_path'] == ''


def test_streaming_structured_tables_disabled_leaves_no_csv(pipeline_stub, tmp_path, manifest):
    """When parser does not save structured tables, streaming audit should not invent CSVs."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    full_text, mda_text = _good_streaming_texts()

    def parse_result(_pdf_path, _save_output):
        source_dir = tmp_path / 'parser_output_no_tables'
        source_dir.mkdir()
        (source_dir / 'full_text.txt').write_text(full_text, encoding='utf-8')
        (source_dir / 'mda_text.txt').write_text(mda_text, encoding='utf-8')
        return {
            'text': full_text,
            'mda_text': mda_text,
            'text_path': str(source_dir / 'full_text.txt'),
            'mda_path': str(source_dir / 'mda_text.txt'),
            'output_dir': str(source_dir),
        }

    pipeline_stub.parser = StubStreamingParser(parse_result)

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=True,
        manifest=manifest,
    )

    audit_dirs = [path for path in Path(pipeline_stub.streaming_audit_output_path).iterdir()]
    assert result is not None
    assert len(audit_dirs) == 1
    assert not list(audit_dirs[0].glob('*.csv'))
    assert Path(result['text_path']).exists()


def test_streaming_retention_applies_after_export_not_per_report(pipeline_stub, tmp_path, monkeypatch):
    """The audit cap must not run mid-batch. Pruning per report deleted earlier
    audit text while the run was still exporting its text_path, so any batch
    larger than the cap shipped a summary pointing at missing directories."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.streaming_max_audit_reports = 2
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    full_text, mda_text = _good_streaming_texts()
    counter = {'value': 0}

    def parse_result(_pdf_path, _save_output):
        idx = counter['value']
        counter['value'] += 1
        source_dir = tmp_path / f'parser_output_{idx}'
        source_dir.mkdir()
        (source_dir / 'full_text.txt').write_text(full_text, encoding='utf-8')
        (source_dir / 'mda_text.txt').write_text(mda_text, encoding='utf-8')
        return {
            'text': full_text,
            'mda_text': mda_text,
            'text_path': str(source_dir / 'full_text.txt'),
            'mda_path': str(source_dir / 'mda_text.txt'),
            'output_dir': str(source_dir),
        }

    pipeline_stub.parser = StubStreamingParser(parse_result)
    monkeypatch.setattr('src.pipeline.create_timestamp', lambda: '20240101_120000')

    exported_text_paths = []
    dangling_at_export = []
    save_results = pipeline_stub.save_results

    def spy_save_results(df):
        paths = df['text_path'].tolist()
        exported_text_paths.extend(paths)
        dangling_at_export.extend(p for p in paths if not Path(p).exists())
        return save_results(df)

    monkeypatch.setattr(pipeline_stub, 'save_results', spy_save_results)

    results = pipeline_stub.run_streaming(
        company_codes=['000001', '000002', '000003'],
        years=[2023],
        report_types=['annual'],
    )

    assert len(results) == 3
    assert len(exported_text_paths) == 3
    assert dangling_at_export == [], \
        f"every exported text_path must resolve at export time: {dangling_at_export}"

    audit_dirs = [path for path in Path(pipeline_stub.streaming_audit_output_path).iterdir()]
    assert len(audit_dirs) == 2, "the cap still applies once the run is done"


def test_streaming_keep_pdf_preserves_file_path_and_pdf(pipeline_stub, tmp_path, manifest):
    """delete_pdf=False should keep the PDF path in results and leave the file on disk."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    full_text, mda_text = _good_streaming_texts()
    pipeline_stub.parser = StubStreamingParser(
        lambda _pdf_path, _save_output: {
            'text': full_text,
            'mda_text': mda_text,
            'text_path': '',
            'mda_path': '',
            'output_dir': '',
        }
    )

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=False,
        save_parsed_text=False,
        manifest=manifest,
    )

    assert result is not None
    assert result['file_path'] == pipeline_stub.downloader.downloaded_paths[0]
    assert Path(result['file_path']).exists()


def _read_manifest(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def _stub_streaming_success(pipeline_stub, tmp_path):
    """Wire the stubs for a streaming run where every report parses cleanly."""
    pipeline_stub.streaming_audit_output_path = str(tmp_path / 'audit_root')
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    full_text, mda_text = _good_streaming_texts()
    pipeline_stub.parser = StubStreamingParser(
        lambda _pdf_path, _save_output: {
            'text': full_text,
            'mda_text': mda_text,
            'text_path': '',
            'mda_path': '',
            'output_dir': '',
        }
    )


def test_process_single_report_records_download_failure(pipeline_stub, tmp_path, manifest):
    """A failed download must be attributable from the ledger, not only from a
    log line a long run has already scrolled past."""
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    pipeline_stub.downloader.download_one_sync = lambda url, save_path: False
    pipeline_stub.parser = StubStreamingParser(lambda _p, _s: {})

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=False,
        manifest=manifest,
    )

    key = StreamingManifest.task_key('000001', 2023, 'annual')
    assert result is None
    assert manifest.entries[key]['status'] == 'download_failed'
    assert _read_manifest(manifest.path)[key]['status'] == 'download_failed'


def test_process_single_report_records_empty_text(pipeline_stub, tmp_path, manifest):
    """A PDF that yields no usable text is a distinct failure from a download
    that never landed, and the two need different fixes."""
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    pipeline_stub.parser = StubStreamingParser(
        lambda _p, _s: {'text': '', 'mda_text': '', 'text_path': '',
                        'mda_path': '', 'output_dir': ''}
    )

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=False,
        manifest=manifest,
    )

    key = StreamingManifest.task_key('000001', 2023, 'annual')
    assert result is None
    assert manifest.entries[key]['status'] == 'no_text'


def test_process_single_report_records_unusable_announcement(pipeline_stub, tmp_path, manifest):
    """An announcement with no downloadable URL used to return None silently."""
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')
    pipeline_stub.downloader.build_download_url = lambda announcement: ('', '')
    pipeline_stub.parser = StubStreamingParser(lambda _p, _s: {})

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=False,
        manifest=manifest,
    )

    key = StreamingManifest.task_key('000001', 2023, 'annual')
    assert result is None
    assert manifest.entries[key]['status'] == 'no_url'
    assert manifest.entries[key]['announcement_title'] == '2023年年度报告_0'


def test_process_single_report_records_exception_with_reason(pipeline_stub, tmp_path, manifest):
    """The ledger carries the exception type and message, so a failure can be
    diagnosed without reproducing the run."""
    pipeline_stub.downloader = StubStreamingDownloader(tmp_path / 'downloads')

    def exploding_parse(_pdf_path, _save_output):
        raise RuntimeError('pdfplumber could not open the file')

    pipeline_stub.parser = StubStreamingParser(exploding_parse)

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=False,
        manifest=manifest,
    )

    key = StreamingManifest.task_key('000001', 2023, 'annual')
    assert result is None
    assert manifest.entries[key]['status'] == 'process_error'
    assert manifest.entries[key]['error'] == (
        'RuntimeError: pdfplumber could not open the file'
    )


def test_process_single_report_records_ok_with_full_result(pipeline_stub, tmp_path, manifest):
    """Storing the whole analysis dict is what lets a resumed run export a
    complete summary without re-downloading anything."""
    _stub_streaming_success(pipeline_stub, tmp_path)

    result = pipeline_stub._process_single_report(
        stock_code='000001',
        year=2023,
        report_type='annual',
        announcement=_streaming_announcement(),
        delete_pdf=True,
        save_parsed_text=False,
        manifest=manifest,
    )

    key = StreamingManifest.task_key('000001', 2023, 'annual')
    assert result is not None
    assert manifest.entries[key]['status'] == 'ok'
    assert manifest.entries[key]['result']['stock_code'] == '000001'

    on_disk = _read_manifest(manifest.path)
    assert on_disk[key]['result']['tone_raw'] == result['tone_raw']
    # The atomic rename must not leave its staging file behind.
    assert not Path(manifest.path + '.tmp').exists()


def test_run_streaming_writes_manifest_with_tally(pipeline_stub, tmp_path):
    """Every task lands in the ledger, including the ones that produced nothing."""
    _stub_streaming_success(pipeline_stub, tmp_path)
    downloader = pipeline_stub.downloader
    downloader.query_announcements = (
        lambda code, year, report_type: [] if code == '000002'
        else [_streaming_announcement(0)]
    )

    results = pipeline_stub.run_streaming(
        company_codes=['000001', '000002'],
        years=[2023],
        report_types=['annual'],
    )

    on_disk = _read_manifest(pipeline_stub.streaming_manifest_path)
    assert on_disk['000001/2023/annual']['status'] == 'ok'
    assert on_disk['000002/2023/annual']['status'] == 'no_announcements'
    assert len(results) == 1


def test_run_streaming_survives_a_failing_announcement_query(pipeline_stub, tmp_path):
    """One transient CNINFO failure must not end the run.

    The tasks are independent, and before the query was wrapped an exception
    here discarded every result already accumulated in memory."""
    _stub_streaming_success(pipeline_stub, tmp_path)
    downloader = pipeline_stub.downloader

    def query(code, year, report_type):
        if code == '000002':
            raise RuntimeError('cninfo returned 502')
        return [_streaming_announcement(0)]

    downloader.query_announcements = query

    results = pipeline_stub.run_streaming(
        company_codes=['000001', '000002', '000003'],
        years=[2023],
        report_types=['annual'],
    )

    on_disk = _read_manifest(pipeline_stub.streaming_manifest_path)
    assert len(results) == 2, "the tasks on either side of the failure still ran"
    assert on_disk['000002/2023/annual']['status'] == 'query_error'
    assert 'cninfo returned 502' in on_disk['000002/2023/annual']['error']


def test_run_streaming_resume_skips_completed_tasks(pipeline_stub, tmp_path):
    """Resume must not re-download finished reports, and must still export them."""
    _stub_streaming_success(pipeline_stub, tmp_path)
    downloader = pipeline_stub.downloader
    downloader.query_announcements = (
        lambda code, year, report_type: [] if code == '000002'
        else [_streaming_announcement(0)]
    )

    pipeline_stub.run_streaming(
        company_codes=['000001', '000002'],
        years=[2023],
        report_types=['annual'],
    )
    downloads_after_first_run = len(downloader.downloaded_paths)
    assert downloads_after_first_run == 1

    resumed = pipeline_stub.run_streaming(
        company_codes=['000001', '000002'],
        years=[2023],
        report_types=['annual'],
        resume=True,
    )

    assert len(downloader.downloaded_paths) == downloads_after_first_run, \
        "the ok task was served from the ledger, not re-downloaded"
    assert len(resumed) == 1, "the summary still covers the skipped report"
    assert resumed.iloc[0]['stock_code'] == '000001'


def test_run_streaming_without_resume_starts_a_fresh_ledger(pipeline_stub, tmp_path):
    """The default must be a clean run: a resubmitted web job against a stale
    ledger would otherwise come back empty and look like a silent no-op."""
    _stub_streaming_success(pipeline_stub, tmp_path)
    downloader = pipeline_stub.downloader

    pipeline_stub.run_streaming(company_codes=['000001'], years=[2023],
                                report_types=['annual'])
    first_run_downloads = len(downloader.downloaded_paths)

    results = pipeline_stub.run_streaming(company_codes=['000001'], years=[2023],
                                          report_types=['annual'])

    assert len(downloader.downloaded_paths) == first_run_downloads * 2
    assert len(results) == 1


def test_streaming_manifest_survives_a_run_that_dies_partway(pipeline_stub, tmp_path):
    """The ledger is written per task, so an interrupted run still says what
    finished. Ctrl-C, a cancel and the watchdog all kill the worker without
    unwinding it, and results otherwise live only in memory until the export."""
    _stub_streaming_success(pipeline_stub, tmp_path)
    original_parse = pipeline_stub.parser.parse_pdf
    calls = {'n': 0}

    def parse_then_die(pdf_path, save_output=True, extract_tables=None):
        calls['n'] += 1
        if calls['n'] == 2:
            raise KeyboardInterrupt
        return original_parse(pdf_path, save_output=save_output,
                              extract_tables=extract_tables)

    pipeline_stub.parser.parse_pdf = parse_then_die

    with pytest.raises(KeyboardInterrupt):
        pipeline_stub.run_streaming(
            company_codes=['000001', '000002'],
            years=[2023],
            report_types=['annual'],
        )

    on_disk = _read_manifest(pipeline_stub.streaming_manifest_path)
    assert on_disk['000001/2023/annual']['status'] == 'ok'
    assert '000002/2023/annual' not in on_disk
    assert not Path(pipeline_stub.streaming_manifest_path + '.tmp').exists()


def test_finish_work_requires_explicit_confirmation(pipeline_stub, tmp_path):
    raw_root = tmp_path / 'raw'
    raw_root.mkdir()
    pdf_path = raw_root / 'report.pdf'
    pdf_path.write_bytes(b'%PDF-1.4')

    pipeline_stub.downloader = SimpleNamespace(download_path=str(raw_root))
    pipeline_stub.parser = SimpleNamespace(output_path=str(tmp_path / 'parsed'))

    with pytest.raises(ValueError):
        pipeline_stub.finish_work()

    assert pdf_path.exists()


def test_finish_work_cleans_cache_but_keeps_audit_and_results(pipeline_stub, tmp_path):
    raw_root = tmp_path / 'raw'
    parsed_root = tmp_path / 'parsed'
    audit_root = parsed_root / 'streaming_audit'
    results_root = tmp_path / 'results'
    intermediate_root = results_root / 'intermediate'

    raw_tmp = raw_root / '_streaming_tmp'
    raw_tmp.mkdir(parents=True)
    raw_pdf = raw_tmp / 'latest.pdf'
    raw_pdf.write_bytes(b'%PDF-1.4')
    (raw_root / '.gitkeep').touch()

    non_audit = parsed_root / 'old_parse'
    non_audit.mkdir(parents=True)
    (non_audit / 'full_text.txt').write_text('old text', encoding='utf-8')
    audit_dir = audit_root / '000001_2023_annual_report'
    audit_dir.mkdir(parents=True)
    audit_text = audit_dir / 'full_text.txt'
    audit_text.write_text('audit text', encoding='utf-8')
    (parsed_root / '.gitkeep').touch()

    intermediate_root.mkdir(parents=True)
    (intermediate_root / 'analysis_results_latest.pkl').write_bytes(b'cache')
    results_root.mkdir(exist_ok=True)
    result_file = results_root / 'master_summary.xlsx'
    result_file.write_bytes(b'result')

    pipeline_stub.downloader = SimpleNamespace(download_path=str(raw_root))
    pipeline_stub.parser = SimpleNamespace(output_path=str(parsed_root))
    pipeline_stub.streaming_audit_output_path = str(audit_root)
    pipeline_stub.intermediate_output_path = str(intermediate_root)

    report = pipeline_stub.finish_work(confirm_work_complete=True)

    assert not raw_tmp.exists()
    assert not non_audit.exists()
    assert not (intermediate_root / 'analysis_results_latest.pkl').exists()
    assert audit_text.exists()
    assert result_file.exists()
    assert raw_root.exists()
    assert parsed_root.exists()
    assert str(audit_root) in report['kept_dirs']


def test_intermediate_frames_can_be_saved_and_loaded(pipeline_stub, monkeypatch):
    """Intermediate snapshots should be recoverable for later reruns."""
    pipeline_stub.save_intermediate = True
    monkeypatch.setattr('src.pipeline.create_timestamp', lambda: '20240101_120002')

    frame = pd.DataFrame([{'stock_code': '000001', 'tone_raw': 0.1}])
    saved_path = pipeline_stub._save_intermediate_frame('analysis_results', frame)
    loaded = pipeline_stub.load_intermediate_frame('analysis_results', 'latest')

    assert saved_path is not None
    assert os.path.exists(saved_path)
    assert loaded.equals(frame)


def test_intermediate_frame_can_be_loaded_from_explicit_path(pipeline_stub, tmp_path):
    """Intermediate recovery should support a specific pickle path."""
    frame = pd.DataFrame([{'stock_code': '000002', 'tone_raw': -0.2}])
    explicit_path = tmp_path / 'custom_analysis_results.pkl'
    frame.to_pickle(explicit_path)

    loaded = pipeline_stub.load_intermediate_frame(
        'analysis_results',
        str(explicit_path)
    )

    assert loaded.equals(frame)


def test_intermediate_frame_missing_file_returns_empty(pipeline_stub, tmp_path):
    """Missing intermediate files should return an empty DataFrame."""
    loaded = pipeline_stub.load_intermediate_frame(
        'analysis_results',
        str(tmp_path / 'missing.pkl')
    )

    assert loaded.empty


def test_run_restore_parse_results_skips_download_and_parse(pipeline_stub, tmp_path):
    """Explicit parse_results restore should bypass download/parse and then analyze."""
    parse_frame = pd.DataFrame([
        {
            'stock_code': '000001',
            'company_name': '平安银行',
            'year': 2023,
            'report_type': 'annual',
            'file_path': 'report.pdf',
            'text': '第三章 管理层讨论与分析\n' + ('经营良好。' * 200),
            'mda_text': '第三章 管理层讨论与分析\n' + ('经营良好。' * 100),
            'download_date': '2024-03-15',
        }
    ])
    parse_path = tmp_path / 'parse_results.pkl'
    parse_frame.to_pickle(parse_path)
    analysis_frame = pd.DataFrame([{'stock_code': '000001', 'tone_raw': 0.1}])

    pipeline_stub.download_phase = Mock(return_value={})
    pipeline_stub.parse_phase = Mock(return_value=pd.DataFrame())
    pipeline_stub.analyze_phase = Mock(return_value=analysis_frame)
    pipeline_stub.save_results = Mock(return_value='out.xlsx')

    result = pipeline_stub.run(
        company_codes=['000001'],
        years=[2023],
        restore_parse_results=str(parse_path),
    )

    pipeline_stub.download_phase.assert_not_called()
    pipeline_stub.parse_phase.assert_not_called()
    pipeline_stub.analyze_phase.assert_called_once()
    pd.testing.assert_frame_equal(
        pipeline_stub.analyze_phase.call_args.args[0],
        parse_frame
    )
    assert result.equals(analysis_frame)


def test_run_restore_analysis_results_skips_earlier_phases(pipeline_stub, tmp_path):
    """Explicit analysis_results restore should bypass download, parse, and analyze."""
    analysis_frame = pd.DataFrame([{'stock_code': '000001', 'tone_raw': 0.2}])
    analysis_path = tmp_path / 'analysis_results.pkl'
    analysis_frame.to_pickle(analysis_path)

    pipeline_stub.download_phase = Mock(return_value={})
    pipeline_stub.parse_phase = Mock(return_value=pd.DataFrame())
    pipeline_stub.analyze_phase = Mock(return_value=pd.DataFrame())
    pipeline_stub.save_results = Mock(return_value='out.xlsx')

    result = pipeline_stub.run(
        company_codes=['000001'],
        years=[2023],
        restore_analysis_results=str(analysis_path),
    )

    pipeline_stub.download_phase.assert_not_called()
    pipeline_stub.parse_phase.assert_not_called()
    pipeline_stub.analyze_phase.assert_not_called()
    assert result.equals(analysis_frame)


def test_run_skip_flags_with_skip_analyze_keep_existing_flow(pipeline_stub, tmp_path):
    """Legacy skip flags should still rebuild/load parsed data before loading latest analysis."""
    parse_frame = pd.DataFrame([{'stock_code': '000001', 'text': '正文'}])
    analysis_frame = pd.DataFrame([{'stock_code': '000001', 'tone_raw': 0.3}])
    Path(pipeline_stub.intermediate_output_path).mkdir(parents=True, exist_ok=True)
    analysis_frame.to_pickle(
        Path(pipeline_stub.intermediate_output_path) / 'analysis_results_latest.pkl'
    )

    metadata = {('000001', 2023, 'annual'): [{'file_path': 'report.pdf'}]}
    pipeline_stub._rebuild_metadata_from_files = Mock(return_value=metadata)
    pipeline_stub._load_parsed_results = Mock(return_value=parse_frame)
    pipeline_stub.download_phase = Mock(return_value={})
    pipeline_stub.parse_phase = Mock(return_value=pd.DataFrame())
    pipeline_stub.analyze_phase = Mock(return_value=pd.DataFrame())
    pipeline_stub.save_results = Mock(return_value='out.xlsx')

    result = pipeline_stub.run(
        company_codes=['000001'],
        years=[2023],
        skip_download=True,
        skip_parse=True,
        skip_analyze=True,
    )

    pipeline_stub._rebuild_metadata_from_files.assert_called_once()
    pipeline_stub._load_parsed_results.assert_called_once_with(metadata)
    pipeline_stub.download_phase.assert_not_called()
    pipeline_stub.parse_phase.assert_not_called()
    pipeline_stub.analyze_phase.assert_not_called()
    assert result.equals(analysis_frame)


def test_run_prunes_parsed_output_only_when_it_parsed(pipeline_stub):
    """Retention is an end-of-run step for runs that parsed. A resume run wrote
    no parsed output, so pruning there would delete the very directories the
    next --skip-parse depends on."""
    pruned = []
    pipeline_stub.parser = SimpleNamespace(
        enforce_output_retention=lambda: pruned.append(True)
    )
    pipeline_stub.download_phase = Mock(return_value={})
    pipeline_stub.parse_phase = Mock(
        return_value=pd.DataFrame([{'stock_code': '000001'}])
    )
    pipeline_stub.analyze_phase = Mock(
        return_value=pd.DataFrame([{'stock_code': '000001'}])
    )
    pipeline_stub.save_results = Mock(return_value='out.xlsx')

    pipeline_stub.run(company_codes=['000001'], years=[2023])

    assert pruned == [True], "a parsing run prunes once, after the export"

    pruned.clear()
    pipeline_stub.parse_phase = Mock(return_value=pd.DataFrame())
    pipeline_stub._rebuild_metadata_from_files = Mock(return_value={})
    pipeline_stub._load_parsed_results = Mock(return_value=pd.DataFrame())

    pipeline_stub.run(company_codes=['000001'], years=[2023],
                      skip_download=True, skip_parse=True)

    assert pruned == [], "a resume run must not prune parsed output"


def test_parse_cli_respects_output_override(monkeypatch, tmp_path):
    """CLI parse command should direct outputs to the requested folder."""
    parser_mock = Mock()
    parser_mock.set_output_path = Mock()
    parser_mock.batch_parse = Mock()

    pipeline_instance = Mock()
    pipeline_instance.parser = parser_mock

    pipeline_constructor = Mock(return_value=pipeline_instance)
    monkeypatch.setattr('src.pipeline.FinancialAnalysisPipeline', pipeline_constructor)
    monkeypatch.setattr('src.pipeline.Path.rglob', Mock(return_value=[tmp_path / 'sample.pdf']))
    monkeypatch.setattr(
        'sys.argv',
        [
            'pipeline.py',
            'parse',
            '--input',
            str(tmp_path),
            '--output',
            str(tmp_path / 'parsed_out'),
        ]
    )

    from src.pipeline import main

    main()

    parser_mock.set_output_path.assert_called_once_with(str(tmp_path / 'parsed_out'))
    parser_mock.batch_parse.assert_called_once_with([str(tmp_path / 'sample.pdf')])
    pipeline_constructor.assert_called_once_with(
        config_path='config.yaml',
        sentiment_dict_path=None
    )


def test_analyze_cli_passes_explicit_sentiment_dict(monkeypatch, tmp_path):
    """CLI should only override the config sentiment dictionary when requested."""
    pipeline_instance = Mock()
    pipeline_instance.run.return_value = pd.DataFrame()

    pipeline_constructor = Mock(return_value=pipeline_instance)
    monkeypatch.setattr('src.pipeline.FinancialAnalysisPipeline', pipeline_constructor)
    monkeypatch.setattr(
        'sys.argv',
        [
            'pipeline.py',
            'analyze',
            '--companies',
            str(tmp_path / 'companies.csv'),
            '--years',
            '2023',
            '--config',
            str(tmp_path / 'custom.yaml'),
            '--sentiment-dict',
            str(tmp_path / 'custom_dict.txt'),
        ]
    )

    from src.pipeline import main

    main()

    pipeline_constructor.assert_called_once_with(
        config_path=str(tmp_path / 'custom.yaml'),
        sentiment_dict_path=str(tmp_path / 'custom_dict.txt')
    )


def test_analyze_cli_passes_intermediate_restore_args(monkeypatch, tmp_path):
    """CLI should expose explicit intermediate restore controls."""
    pipeline_instance = Mock()
    pipeline_instance.run.return_value = pd.DataFrame()

    pipeline_constructor = Mock(return_value=pipeline_instance)
    monkeypatch.setattr('src.pipeline.FinancialAnalysisPipeline', pipeline_constructor)
    monkeypatch.setattr(
        'sys.argv',
        [
            'pipeline.py',
            'analyze',
            '--companies',
            str(tmp_path / 'companies.csv'),
            '--years',
            '2023',
            '--restore-parse-results',
            str(tmp_path / 'parse_results.pkl'),
            '--restore-analysis-results',
        ]
    )

    from src.pipeline import main

    main()

    pipeline_instance.run.assert_called_once_with(
        company_csv=str(tmp_path / 'companies.csv'),
        years=[2023],
        report_types=['annual'],
        financial_data_csv=None,
        skip_download=False,
        skip_parse=False,
        skip_analyze=False,
        restore_parse_results=str(tmp_path / 'parse_results.pkl'),
        restore_analysis_results='latest'
    )


def test_finish_work_cli_requires_confirmation(monkeypatch):
    pipeline_instance = Mock()

    pipeline_constructor = Mock(return_value=pipeline_instance)
    monkeypatch.setattr('src.pipeline.FinancialAnalysisPipeline', pipeline_constructor)
    monkeypatch.setattr('sys.argv', ['pipeline.py', 'finish-work'])

    from src.pipeline import main

    with pytest.raises(SystemExit):
        main()

    pipeline_instance.finish_work.assert_not_called()


def test_finish_work_cli_passes_confirmation_and_dry_run(monkeypatch):
    pipeline_instance = Mock()
    pipeline_instance.finish_work.return_value = {
        'removed_files': ['data/raw/report.pdf'],
        'removed_dirs': [],
        'kept_dirs': ['data/parsed/streaming_audit'],
    }

    pipeline_constructor = Mock(return_value=pipeline_instance)
    monkeypatch.setattr('src.pipeline.FinancialAnalysisPipeline', pipeline_constructor)
    monkeypatch.setattr(
        'sys.argv',
        ['pipeline.py', 'finish-work', '--confirm-work-complete', '--dry-run']
    )

    from src.pipeline import main

    main()

    pipeline_instance.finish_work.assert_called_once_with(
        confirm_work_complete=True,
        dry_run=True
    )
